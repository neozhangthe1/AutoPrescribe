from fastapi import FastAPI, HTTPException, Response
from fastapi.middleware.cors import CORSMiddleware
from fastapi.middleware.gzip import GZipMiddleware
from starlette.middleware.base import BaseHTTPMiddleware
from pydantic import BaseModel
from typing import Dict, List
from datetime import datetime
import uuid
import time
from functools import lru_cache

app = FastAPI()

# Add Gzip compression
app.add_middleware(GZipMiddleware, minimum_size=1000)

# Add response timing middleware
class TimingMiddleware(BaseHTTPMiddleware):
    async def dispatch(self, request, call_next):
        start_time = time.time()
        response = await call_next(request)
        process_time = time.time() - start_time
        response.headers["X-Process-Time"] = str(process_time)
        return response

app.add_middleware(TimingMiddleware)

# Add CORS middleware
app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "https://word-editor-app-kgfnbyhb.devinapps.com",
        "http://localhost:5173"
    ],
    allow_credentials=True,
    allow_methods=["GET", "POST", "PUT", "DELETE"],
    allow_headers=["*"]
)

# In-memory document storage
documents: Dict[str, dict] = {}

class Document(BaseModel):
    content: str

@app.get("/")
def root():
    return {"message": "Hello World"}

@app.get("/healthz")
def healthcheck():
    return {"status": "ok"}

@app.post("/documents")
def create_document(title: str):
    if not title or len(title.strip()) == 0:
        raise HTTPException(status_code=400, detail="Title cannot be empty")
    doc_id = str(uuid.uuid4())
    document = {
        "id": doc_id,
        "title": title.strip(),
        "content": "",
        "created_at": datetime.utcnow().isoformat(),
        "updated_at": datetime.utcnow().isoformat()
    }
    documents[doc_id] = document
    return document

@lru_cache(maxsize=1)
def get_cached_documents() -> List[dict]:
    return list(documents.values())

@app.get("/documents")
def list_documents(response: Response):
    # Set cache headers with ETag
    documents_list = get_cached_documents()
    etag = f'W/"{hash(str(documents_list))}"'
    response.headers["ETag"] = etag
    response.headers["Cache-Control"] = "max-age=30"
    return documents_list

@app.post("/documents")
def create_document(title: str):
    if not title or len(title.strip()) == 0:
        raise HTTPException(status_code=400, detail="Title cannot be empty")
    doc_id = str(uuid.uuid4())
    document = {
        "id": doc_id,
        "title": title.strip(),
        "content": "",
        "created_at": datetime.utcnow().isoformat(),
        "updated_at": datetime.utcnow().isoformat()
    }
    documents[doc_id] = document
    # Invalidate cache
    get_cached_documents.cache_clear()
    return document

@app.get("/documents/{doc_id}")
def get_document(doc_id: str):
    if doc_id not in documents:
        raise HTTPException(status_code=404, detail="Document not found")
    return documents[doc_id]

@app.put("/documents/{doc_id}")
def update_document(doc_id: str, document: Document):
    if doc_id not in documents:
        raise HTTPException(status_code=404, detail="Document not found")
    documents[doc_id]["content"] = document.content
    documents[doc_id]["updated_at"] = datetime.utcnow().isoformat()
    return documents[doc_id]

@app.delete("/documents/{doc_id}")
def delete_document(doc_id: str):
    if doc_id not in documents:
        raise HTTPException(status_code=404, detail="Document not found")
    del documents[doc_id]
    return {"status": "success", "message": "Document deleted"}

from fastapi.responses import StreamingResponse
from io import BytesIO
from docx import Document as DocxDocument
from bs4 import BeautifulSoup
import re

@app.get("/documents/{doc_id}/download")
async def download_document(doc_id: str):
    if doc_id not in documents:
        raise HTTPException(status_code=404, detail="Document not found")
    
    doc = documents[doc_id]
    docx = DocxDocument()
    docx.add_heading(doc["title"], 0)
    
    # Convert HTML to plain text with basic formatting
    if doc["content"]:
        soup = BeautifulSoup(doc["content"], "html.parser")
        for element in soup.find_all(True):
            if element.name == "h1":
                docx.add_heading(element.get_text(), 1)
            elif element.name == "h2":
                docx.add_heading(element.get_text(), 2)
            elif element.name == "p":
                docx.add_paragraph(element.get_text())
            elif element.name == "ul":
                for li in element.find_all("li"):
                    docx.add_paragraph(li.get_text(), style="List Bullet")
    
    # Save to BytesIO
    docx_file = BytesIO()
    docx.save(docx_file)
    docx_file.seek(0)
    
    filename = re.sub(r'[^\w\s-]', '', doc["title"]).strip().lower()
    filename = re.sub(r'[-\s]+', '-', filename)
    
    return StreamingResponse(
        docx_file,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{filename}.docx"'}
    )

@app.get("/api/docs")
async def get_docs():
    """
    Get API documentation information
    """
    return {
        "docs": "Visit /docs or /redoc for API documentation",
        "endpoints": {
            "documents": {
                "list": {"method": "GET", "path": "/documents"},
                "create": {"method": "POST", "path": "/documents?title=string"},
                "get": {"method": "GET", "path": "/documents/{doc_id}"},
                "update": {"method": "PUT", "path": "/documents/{doc_id}"},
                "delete": {"method": "DELETE", "path": "/documents/{doc_id}"},
                "download": {"method": "GET", "path": "/documents/{doc_id}/download"}
            },
            "health": {"method": "GET", "path": "/healthz"}
        }
    }
